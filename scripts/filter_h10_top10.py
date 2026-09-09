#!/usr/bin/env python3
"""Validate an H10 natural-keyword export and return all rank 1-10 phrases."""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import warnings
from pathlib import Path
from typing import Any, Iterable


KEYWORD_ALIASES = {
    "keyword",
    "keywords",
    "keywordphrase",
    "searchphrase",
    "searchterm",
    "关键词",
    "关键词短语",
    "关键词词组",
    "关键字",
    "�ؼ��ʴ���",  # Mojibake seen in some H10 exports.
}
RANK_ALIASES = {
    "organicrank",
    "organicposition",
    "naturalrank",
    "naturalposition",
    "自然排名",
    "自然位",
    "��Ȼ����",  # Mojibake seen in some H10 exports.
}
VOLUME_ALIASES = {
    "searchvolume",
    "monthlysearchvolume",
    "搜索量",
    "月搜索量",
    "������",  # Mojibake seen in some H10 exports.
}


def normalize_header(value: Any) -> str:
    text = "" if value is None else str(value).strip().lower()
    return re.sub(r"[\s_\-—–:/\\()（）]+", "", text)


def parse_rank(value: Any) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return int(number) if number.is_integer() and number > 0 else None
    text = str(value).strip().replace(",", "")
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        number = float(text)
        return int(number) if number.is_integer() and number > 0 else None
    return None


def parse_volume(value: Any) -> int | float | str | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value).strip()
    numeric = text.replace(",", "")
    if re.fullmatch(r"-?\d+(?:\.\d+)?", numeric):
        number = float(numeric)
        return int(number) if number.is_integer() else number
    return text


def read_delimited(path: Path) -> list[tuple[str, list[list[Any]]]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                rows = [list(row) for row in csv.reader(handle, delimiter=delimiter)]
            return [(path.stem, rows)]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"Unable to decode delimited file: {last_error}")


def read_excel(path: Path, requested_sheet: str | None) -> list[tuple[str, list[list[Any]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Excel input requires openpyxl. Load the bundled workspace dependencies "
            "and run this script with their Python executable."
        ) from exc

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = load_workbook(path, read_only=True, data_only=True)

    if requested_sheet and requested_sheet not in workbook.sheetnames:
        raise ValueError(
            f"Worksheet {requested_sheet!r} was not found. Available: {workbook.sheetnames}"
        )

    names = [requested_sheet] if requested_sheet else workbook.sheetnames
    result: list[tuple[str, list[list[Any]]]] = []
    for name in names:
        worksheet = workbook[name]
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        result.append((name, rows))
    workbook.close()
    return result


def find_table(
    sheets: Iterable[tuple[str, list[list[Any]]]],
) -> tuple[str, list[list[Any]], int, int, int, int | None]:
    inspected_headers: list[str] = []
    for sheet_name, rows in sheets:
        for row_index, row in enumerate(rows[:20]):
            normalized = [normalize_header(cell) for cell in row]
            keyword_index = next(
                (i for i, value in enumerate(normalized) if value in KEYWORD_ALIASES),
                None,
            )
            rank_index = next(
                (i for i, value in enumerate(normalized) if value in RANK_ALIASES),
                None,
            )
            volume_index = next(
                (i for i, value in enumerate(normalized) if value in VOLUME_ALIASES),
                None,
            )
            if keyword_index is not None and rank_index is not None:
                return (
                    sheet_name,
                    rows,
                    row_index,
                    keyword_index,
                    rank_index,
                    volume_index,
                )
            if any(normalized):
                inspected_headers.append(f"{sheet_name}!{row_index + 1}: {row}")

    preview = "\n".join(inspected_headers[:8])
    raise ValueError(
        "Could not identify both the complete-keyword and natural/organic-rank columns "
        "within the first 20 rows of any worksheet. Rename the columns to clear labels "
        "such as 'Keyword Phrase' and 'Organic Rank' or '关键词短语' and '自然排名'."
        + (f"\nInspected header candidates:\n{preview}" if preview else "")
    )


def extract_asin(path: Path) -> str | None:
    match = re.search(r"(?i)(?<![A-Z0-9])(B0[A-Z0-9]{8})(?![A-Z0-9])", path.stem)
    return match.group(1).upper() if match else None


def collect(path: Path, requested_sheet: str | None, max_rank: int) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        sheets = read_excel(path, requested_sheet)
    elif suffix in {".csv", ".tsv"}:
        sheets = read_delimited(path)
    else:
        raise ValueError("Supported input formats are .xlsx, .xlsm, .csv, and .tsv")

    sheet_name, rows, header_index, keyword_index, rank_index, volume_index = find_table(sheets)
    header = rows[header_index]
    data_rows = rows[header_index + 1 :]
    valid_rank_rows = 0
    selected: dict[str, dict[str, Any]] = {}

    for source_row, row in enumerate(data_rows, start=header_index + 2):
        rank_value = row[rank_index] if rank_index < len(row) else None
        rank = parse_rank(rank_value)
        if rank is None:
            continue
        valid_rank_rows += 1
        if rank > max_rank:
            continue

        keyword_value = row[keyword_index] if keyword_index < len(row) else None
        keyword = "" if keyword_value is None else str(keyword_value).strip()
        if not keyword:
            continue

        volume_value = row[volume_index] if volume_index is not None and volume_index < len(row) else None
        item = {
            "keyword": keyword,
            "natural_rank": rank,
            "search_volume": parse_volume(volume_value),
            "source_row": source_row,
        }
        key = keyword.casefold()
        existing = selected.get(key)
        if existing is None or (rank, source_row) < (
            existing["natural_rank"],
            existing["source_row"],
        ):
            selected[key] = item

    keywords = sorted(
        selected.values(), key=lambda item: (item["natural_rank"], item["keyword"].casefold())
    )
    return {
        "source_file": str(path.resolve()),
        "asin_from_filename": extract_asin(path),
        "worksheet": sheet_name,
        "header_row": header_index + 1,
        "detected_columns": {
            "keyword": str(header[keyword_index]),
            "natural_rank": str(header[rank_index]),
            "search_volume": (
                str(header[volume_index]) if volume_index is not None else None
            ),
        },
        "rank_scope": f"1-{max_rank}",
        "total_data_rows": len(data_rows),
        "valid_natural_rank_rows": valid_rank_rows,
        "inspected_keyword_count": len(keywords),
        "keywords": keywords,
    }


def write_output(result: dict[str, Any], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.suffix.lower() == ".json":
        output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        return
    if output.suffix.lower() == ".csv":
        with output.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=["keyword", "natural_rank", "search_volume", "source_row"],
            )
            writer.writeheader()
            writer.writerows(result["keywords"])
        return
    raise ValueError("Output must end in .json or .csv")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Extract every distinct H10 natural keyword ranked 1-10."
    )
    parser.add_argument("input", type=Path, help="H10 .xlsx, .xlsm, .csv, or .tsv file")
    parser.add_argument("--sheet", help="Worksheet name when automatic detection is unsuitable")
    parser.add_argument(
        "--max-rank",
        type=int,
        default=10,
        help="Inclusive natural-rank ceiling (default: 10)",
    )
    parser.add_argument("--output", type=Path, help="Optional .json or .csv output path")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    try:
        if not args.input.is_file():
            raise ValueError(f"Input file does not exist: {args.input}")
        if args.max_rank < 1:
            raise ValueError("--max-rank must be at least 1")
        result = collect(args.input, args.sheet, args.max_rank)
        if args.output:
            write_output(result, args.output)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (OSError, RuntimeError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
